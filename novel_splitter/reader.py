"""
reader.py
---------
원본 소설 파일(docx, txt)을 읽어 내부 표현(Document)으로 변환한다.

내부 표현은 '블록(Block)들의 리스트' 이다.
- docx : 문단(ParagraphBlock)과 표(TableBlock)를 '문서에 나타난 순서대로' 읽는다.
         (표 안의 텍스트도 누락 없이 포함된다)
- txt  : 각 줄(line)이 하나의 ParagraphBlock

이 통일된 표현 덕분에 splitter/counter/writer 가 파일 형식이나
표 유무에 관계없이 동일한 방식으로 동작할 수 있다.

txt 인코딩은 UTF-8, UTF-8 BOM, CP949 를 자동 감지한다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import List

from .blocks import Block, ParagraphBlock, TableBlock
from .utils import get_file_extension


@dataclass
class Document:
    """읽어들인 문서의 내부 표현."""

    blocks: List[Block] = field(default_factory=list)  # 문단/표 블록 목록(순서 보존)
    file_format: str = "txt"  # 'docx' 또는 'txt'
    encoding: str = "utf-8"   # txt 원본 인코딩 (docx는 무시)
    source_path: str = ""     # 원본 파일 경로

    @property
    def paragraphs(self) -> List[str]:
        """
        하위 호환용: 각 블록의 대표 텍스트 목록.
        표는 대표 텍스트가 비어 있으므로 분량 계산에는 blocks 를 사용해야 한다.
        """
        return [b.text for b in self.blocks]

    @property
    def full_text(self) -> str:
        """전체 텍스트(계산용)를 이어 붙여 반환한다."""
        return "\n".join(b.count_text() for b in self.blocks)


# txt 인코딩 자동 감지 시도 순서
# 1) utf-8-sig : BOM이 있으면 처리하고 BOM 없는 UTF-8도 읽음
# 2) cp949     : 한국어 Windows 기본 인코딩(EUC-KR 상위호환)
_TXT_ENCODINGS = ("utf-8-sig", "cp949")


class Reader:
    """파일 형식에 맞춰 문서를 읽어 Document 로 반환하는 클래스."""

    def read(self, path: str) -> Document:
        """확장자에 따라 적절한 리더를 호출한다."""
        ext = get_file_extension(path)
        if ext == ".docx":
            return self._read_docx(path)
        elif ext == ".txt":
            return self._read_txt(path)
        raise ValueError(f"지원하지 않는 파일 형식입니다: {ext}")

    # ------------------------------------------------------------------ #
    # 내부 구현
    # ------------------------------------------------------------------ #
    def _read_docx(self, path: str) -> Document:
        """
        docx 를 문단·표 블록으로 '순서대로' 읽는다.
        python-docx 의 doc.paragraphs 는 표 내용을 포함하지 않으므로,
        문서 본문(body)의 자식 요소를 직접 순회한다.
        """
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "docx 파일을 읽으려면 python-docx 가 필요합니다. "
                "'pip install python-docx' 로 설치하세요."
            ) from exc

        docx_doc = DocxDocument(path)
        blocks: List[Block] = []
        for item in self._iter_block_items(docx_doc):
            blocks.append(item)

        return Document(
            blocks=blocks,
            file_format="docx",
            encoding="utf-8",
            source_path=path,
        )

    @staticmethod
    def _iter_block_items(docx_doc):
        """
        docx 본문을 순서대로 순회하며 ParagraphBlock / TableBlock 을 생성한다.
        문단(<w:p>)과 표(<w:tbl>)가 원본에 나타난 순서를 그대로 보존한다.
        """
        from docx.document import Document as _DocClass
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table as _Table
        from docx.text.paragraph import Paragraph as _Paragraph

        parent_elm = docx_doc.element.body
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                para = _Paragraph(child, docx_doc)
                yield ParagraphBlock(text=para.text)
            elif isinstance(child, CT_Tbl):
                table = _Table(child, docx_doc)
                rows = [
                    [cell.text for cell in row.cells] for row in table.rows
                ]
                yield TableBlock(rows=rows)

    def _read_txt(self, path: str) -> Document:
        """
        txt 파일을 인코딩 자동 감지하여 읽는다.
        UTF-8(BOM 포함) -> CP949 순으로 시도한다.
        각 줄은 ParagraphBlock 이 된다.
        """
        raw = self._read_bytes(path)

        text = None
        used_encoding = "utf-8"
        for enc in _TXT_ENCODINGS:
            try:
                text = raw.decode(enc)
                # utf-8-sig 로 성공하면 실제 인코딩 이름을 정리
                used_encoding = "utf-8" if enc == "utf-8-sig" else enc
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if text is None:
            # 마지막 수단: 손실 허용(replace) UTF-8 디코딩
            text = raw.decode("utf-8", errors="replace")
            used_encoding = "utf-8"

        # 개행 문자를 통일한 뒤 줄 단위로 분리한다.
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        blocks: List[Block] = [ParagraphBlock(text=line) for line in text.split("\n")]

        return Document(
            blocks=blocks,
            file_format="txt",
            encoding=used_encoding,
            source_path=path,
        )

    @staticmethod
    def _read_bytes(path: str) -> bytes:
        """파일을 바이트로 읽는다."""
        with open(path, "rb") as f:
            return f.read()
