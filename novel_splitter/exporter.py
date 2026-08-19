"""
exporter.py
-----------
분권 결과(FileResult 목록)를 Excel(.xlsx) 파일로 내보낸다.

결과 표에 표시되는 항목(번호/파일명/공백포함·제외 글자수/단어수/줄수)에
원본 파일명을 추가로 담고, 맨 아래에 총합 행을 넣는다.

openpyxl 을 사용한다(프로젝트 의존성에 포함되어 있음).
"""

from __future__ import annotations

import os
from typing import List

from .counter import Counts

# 결과 표의 열 정의: (헤더, 숫자 여부)
_COLUMNS = [
    ("번호", False),
    ("원본파일", False),
    ("파일명", False),
    ("공백포함 글자수", True),
    ("공백제외 글자수", True),
    ("단어수", True),
    ("줄수", True),
]


class ExcelExporter:
    """분권 결과를 .xlsx 로 저장하는 클래스."""

    def export(self, results: List["FileResult"], path: str) -> None:
        """
        results 를 Excel 파일로 저장한다.

        Parameters
        ----------
        results : Processor.process_file 가 반환한 FileResult 들의 목록
        path : 저장할 .xlsx 경로
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "Excel 내보내기에는 openpyxl 이 필요합니다. "
                "'pip install openpyxl' 로 설치하세요."
            ) from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "분권 결과"

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="D9E1F2")
        total_font = Font(bold=True)
        total_fill = PatternFill("solid", fgColor="FCE4D6")
        right = Alignment(horizontal="right")
        center = Alignment(horizontal="center")

        # 헤더 행
        for col, (title, _is_num) in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # 데이터 행
        row_idx = 2
        seq = 0
        grand_total = Counts()
        for file_result in results:
            source_name = os.path.basename(file_result.source_path)
            for chunk in file_result.chunks:
                seq += 1
                filename = chunk.filename + (" (건너뜀)" if chunk.skipped else "")
                values = [
                    seq,
                    source_name,
                    filename,
                    chunk.counts.chars_with_spaces,
                    chunk.counts.chars_without_spaces,
                    chunk.counts.words,
                    chunk.counts.lines,
                ]
                for col, ((_title, is_num), value) in enumerate(
                    zip(_COLUMNS, values), start=1
                ):
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    if is_num:
                        cell.number_format = "#,##0"
                        cell.alignment = right
                grand_total = grand_total + chunk.counts
                row_idx += 1

        # 총합 행
        total_values = [
            "총합",
            "",
            f"분권 {seq}개",
            grand_total.chars_with_spaces,
            grand_total.chars_without_spaces,
            grand_total.words,
            grand_total.lines,
        ]
        for col, ((_title, is_num), value) in enumerate(
            zip(_COLUMNS, total_values), start=1
        ):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = total_font
            cell.fill = total_fill
            if is_num:
                cell.number_format = "#,##0"
                cell.alignment = right

        # 열 너비 자동 조정(대략적인 최대 길이 기반)
        self._autofit_columns(ws, get_column_letter, ncols=len(_COLUMNS), last_row=row_idx)

        # 헤더 고정(스크롤해도 헤더가 보이도록)
        ws.freeze_panes = "A2"

        wb.save(path)

    # ------------------------------------------------------------------ #
    # 파일별 분량 내보내기 (분권된 파일 분량 파악용)
    # ------------------------------------------------------------------ #
    def export_file_analysis(self, items, path: str) -> None:
        """
        파일별 분량을 Excel 로 저장한다.

        items : (파일명, Counts) 튜플의 리스트
        맨 아래에 총합/평균/최대/최소 요약 행을 넣는다.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover
            raise ImportError("Excel 내보내기에는 openpyxl 이 필요합니다.") from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "파일 분량"

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="D9E1F2")
        sum_font = Font(bold=True)
        sum_fill = PatternFill("solid", fgColor="FCE4D6")
        right = Alignment(horizontal="right")
        center = Alignment(horizontal="center")

        headers = ["번호", "파일명", "공백포함 글자수", "공백제외 글자수", "단어수", "줄수"]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        row_idx = 2
        chars_with = []
        chars_without = []
        words = []
        lines = []
        for i, (name, counts) in enumerate(items, start=1):
            values = [
                i, name,
                counts.chars_with_spaces, counts.chars_without_spaces,
                counts.words, counts.lines,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if col >= 3:
                    cell.number_format = "#,##0"
                    cell.alignment = right
            chars_with.append(counts.chars_with_spaces)
            chars_without.append(counts.chars_without_spaces)
            words.append(counts.words)
            lines.append(counts.lines)
            row_idx += 1

        # 요약 행: 총합 / 평균 / 최대 / 최소
        n = max(len(items), 1)

        def summary_row(label, fn):
            nonlocal row_idx
            vals = [label, "",
                    fn(chars_with), fn(chars_without), fn(words), fn(lines)]
            for col, value in enumerate(vals, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = sum_font
                cell.fill = sum_fill
                if col >= 3:
                    cell.number_format = "#,##0"
                    cell.alignment = right
            row_idx += 1

        if items:
            summary_row("총합", lambda a: sum(a))
            summary_row("평균", lambda a: sum(a) // n)
            summary_row("최대", lambda a: max(a))
            summary_row("최소", lambda a: min(a))

        self._autofit_columns(ws, get_column_letter, ncols=len(headers), last_row=row_idx)
        ws.freeze_panes = "A2"
        wb.save(path)

    # ------------------------------------------------------------------ #
    # 플랫폼 회차 목록 내보내기 (분권 전 확인용)
    # ------------------------------------------------------------------ #
    def export_platform(self, result: "PlatformResult", path: str) -> None:
        """
        플랫폼에서 가져온 회차 목록(번호/제목)을 Excel 로 저장한다.
        분권 전에 플랫폼 연재 구성을 미리 확인하는 용도.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "Excel 내보내기에는 openpyxl 이 필요합니다."
            ) from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "플랫폼 회차목록"

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="D9E1F2")
        center = Alignment(horizontal="center")

        # 1행: 작품 정보(참고용)
        ws.cell(row=1, column=1, value="작품").font = header_font
        ws.cell(row=1, column=2, value=result.work_title or "(제목 미확인)")
        ws.cell(row=2, column=1, value="총 회차").font = header_font
        ws.cell(row=2, column=2, value=result.total_count)
        ws.cell(row=3, column=1, value="출처 URL").font = header_font
        ws.cell(row=3, column=2, value=result.source_url)

        # 5행: 표 헤더
        header_row = 5
        for col, title in enumerate(["번호", "회차 제목"], start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # 회차 데이터
        row_idx = header_row + 1
        for ep in result.episodes:
            ws.cell(row=row_idx, column=1, value=ep.no)
            ws.cell(row=row_idx, column=2, value=ep.title)
            row_idx += 1

        self._autofit_columns(ws, get_column_letter, ncols=2, last_row=row_idx)
        ws.freeze_panes = f"A{header_row + 1}"
        wb.save(path)

    @staticmethod
    def _autofit_columns(ws, get_column_letter, ncols: int, last_row: int) -> None:
        """각 열의 내용 길이에 맞춰 너비를 대략 조정한다."""
        for col in range(1, ncols + 1):
            max_len = 0
            for row in range(1, last_row + 1):
                value = ws.cell(row=row, column=col).value
                if value is None:
                    continue
                # 한글은 폭이 넓으므로 약간 가중
                text = str(value)
                width = sum(2 if ord(ch) > 0x1100 else 1 for ch in text)
                max_len = max(max_len, width)
            ws.column_dimensions[get_column_letter(col)].width = min(
                max(max_len + 2, 8), 60
            )
