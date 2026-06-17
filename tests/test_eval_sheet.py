"""평가표(.xlsx) 생성 엔진 검증 — 데이터 기입 + 수식/드롭다운 보존 + 엣지케이스."""

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from mtpe.evaluation.sheet import generate_files  # noqa: E402


def _data(**over):
    d = {
        "work_id": "WORK01", "language_pair": "ko-en", "prompt_version": "v1.2",
        "extraction_date": "2026-06-15",
        "source_sentences": ["가.", "나.", "다."],
        "mt_sentences": ["a.", "b.", "c."],
        "scene_breakdown": [
            {"scene_id": "장면-01", "scene_description": "첫 장면",
             "start_sentence_index": 0, "end_sentence_index": 2},
        ],
        "output_mode": "single",
    }
    d.update(over)
    return d


def test_single_fills_and_preserves():
    res = generate_files(_data())
    assert len(res.files) == 1
    name, payload = res.files[0]
    assert name == "WORK01_v1.2_평가표.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(payload))
    ws = wb["③ 평가 입력"]
    # 데이터 기입
    assert ws["A3"].value == "WORK01-001"
    assert ws["B3"].value == "가."
    assert ws["C3"].value == "a."
    # 카테고리는 비어 있어야(평가자 입력)
    assert ws["E3"].value is None
    # 집계 수식·드롭다운 보존
    assert str(ws["S3"].value).startswith("=IF(")
    assert any("E3:R349" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    # ④ 장면
    qws = wb["④ 정성 평가"]
    assert qws["A3"].value == "장면-01"
    assert str(qws["J3"].value).startswith("=")
    # ⑤ ⑥
    assert wb["⑤ 결과 대시보드"]["B4"].value == "v1.2"
    assert wb["⑥ 버전 비교"]["C4"].value == "v1.2 (현재)"


def test_method_a_three_files():
    res = generate_files(_data(output_mode="method_a"))
    names = [f[0] for f in res.files]
    assert names == [
        "WORK01_v1.2_평가표_A용.xlsx",
        "WORK01_v1.2_평가표_B용.xlsx",
        "WORK01_v1.2_평가표_최종본.xlsx",
    ]


def test_length_mismatch_errors():
    with pytest.raises(ValueError):
        generate_files(_data(mt_sentences=["a.", "b."]))


def test_scene_index_out_of_range_errors():
    with pytest.raises(ValueError):
        generate_files(_data(scene_breakdown=[
            {"scene_id": "s", "scene_description": "d",
             "start_sentence_index": 0, "end_sentence_index": 99}]))


def test_filename_sanitized():
    res = generate_files(_data(work_id="WO/RK:01"))
    assert "/" not in res.files[0][0] and ":" not in res.files[0][0]


def test_ko_ja_layout_human_and_mt_columns():
    """한일은 사람번역 C, MT D 위치 + F~R 카테고리."""
    res = generate_files(_data(
        language_pair="ko-ja", prompt_version="v1.0",
        source_sentences=["원문1.", "원문2."],
        human_translation=["人間訳1。", "人間訳2。"],
        mt_sentences=["MT訳1。", "MT訳2。"],
        scene_breakdown=[{"scene_id": "場面-01", "scene_description": "첫 장면",
                          "start_sentence_index": 0, "end_sentence_index": 1}],
    ))
    wb = openpyxl.load_workbook(io.BytesIO(res.files[0][1]))
    ws = wb["③ 평가 입력"]
    assert ws["B3"].value == "원문1."        # 원문
    assert ws["C3"].value == "人間訳1。"      # 사람 번역본
    assert ws["D3"].value == "MT訳1。"       # MT본
    assert str(ws["S3"].value).startswith("=IF(")  # 집계 수식 보존
    assert any("F3:R203" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
