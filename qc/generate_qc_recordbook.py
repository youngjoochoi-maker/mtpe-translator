"""
MTPE QC 기록부 생성기

사용법:
    pip install openpyxl
    python generate_qc_recordbook.py
  → 같은 폴더에 MTPE_QC_기록부.xlsx 생성

프로덕트가 바뀌면 이 스크립트의 rounds / chk 부분만 고쳐 다시 실행하면
기록부가 새로 만들어집니다. (예: 평가표는 현재 1부 생성 기준)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT="맑은 고딕"
NAVY="1F4E5F"; TEAL="2E7D8A"; LGRAY="EEF3F5"; SECT="D7E6EA"
CORE="FFF3D6"; FILLIN="FFFBEA"
thin=Side(style="thin",color="BFCBD0")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def st(c,*,bold=False,size=10,color="000000",bg=None,align="left",wrap=False,bd=True,top=False):
    c.font=Font(name=FONT,bold=bold,size=size,color=color)
    c.alignment=Alignment(horizontal=align,vertical="top" if top else "center",wrap_text=wrap)
    if bg:c.fill=PatternFill("solid",fgColor=bg)
    if bd:c.border=border
    return c
def hd(c,t,bg=NAVY):
    c.value=t; st(c,bold=True,color="FFFFFF",bg=bg,align="center",wrap=True)

wb=Workbook()

# ===== 1. 시작하기 =====
g=wb.active; g.title="시작하기"; g.sheet_view.showGridLines=False
g.column_dimensions["A"].width=3
g.column_dimensions["B"].width=22
g.column_dimensions["C"].width=86
g.merge_cells("B2:C2"); st(g["B2"],bold=True,size=18,color="FFFFFF",bg=NAVY)
g["B2"]="  MTPE 품질 점검(QC) — 처음 하는 분을 위한 기록장"
g.row_dimensions[2].height=34

def block(row,title):
    g.merge_cells(f"B{row}:C{row}")
    st(g[f"B{row}"],bold=True,size=11,color="FFFFFF",bg=TEAL); g[f"B{row}"]="  "+title
    return row+1
def line(row,k,v,kbg=LGRAY):
    st(g[f"B{row}"],bold=True,bg=kbg,align="left",wrap=True); g[f"B{row}"]=k
    st(g[f"C{row}"],align="left",wrap=True); g[f"C{row}"]=v
    g.row_dimensions[row].height=max(16,18*(len(v)//52+1))
    return row+1

r=4
r=block(r,"① 이 파일은 무엇인가요?")
r=line(r,"한 줄 설명","번역 프로그램(MTPE)이 제대로 작동하는지 직접 눌러보며 확인하고, 그 결과를 적어 두는 기록장입니다.")
r=line(r,"QC가 뭔가요?","Quality Check, 즉 '품질 점검'. 만든 사람 말고 실제로 써 보면서 '여기까지 잘 되네 / 여기서 막히네'를 찾아내는 일이에요. 개발 지식이 없어도 화면을 눌러보며 할 수 있습니다.")
r=line(r,"이렇게 진행해요","① 돈 안 드는 '모의(연습) 모드'로 전체 흐름 먼저 → ② 진짜 모델(Gemini)로 번역 품질 확인 → ③ 일부러 이상하게 해보고 껐다 켜도 멀쩡한지 확인. 이 순서대로 '단계별 따라하기' 시트를 위에서 아래로 하면 됩니다.")

r+=1
r=block(r,"② 어느 시트부터 보나요?")
r=line(r,"단계별 따라하기 ⭐","여기가 본 작업장. 순서대로 '해볼 것'을 따라 하고, '이렇게 나오면 정상'과 비교해 결과를 고르면 됩니다.")
r=line(r,"세부 점검표","더 자세히 항목별로 확인하고 싶을 때 보는 곳(선택). 안 봐도 되지만, 꼼꼼히 보려면 유용해요.")
r=line(r,"문제 적기","하다가 이상하거나 막히면 여기에 적습니다. 화면을 캡처해 파일명도 함께 적어두면 좋아요.")
r=line(r,"종합 결과","위 결과들을 모아 '출시해도 되는지'를 자동으로 알려줍니다. 직접 계산할 필요 없어요.")

r+=1
r=block(r,"③ 결과는 어떻게 적나요?")
r=line(r,"칸을 누르면","연한 노란 칸을 누르면 작은 화살표(▼)가 생겨요. 눌러서 골라 넣으면 됩니다. 직접 타이핑 안 해도 돼요.")
r=line(r,"고르는 보기","✅ 잘 됨 / ❌ 안 됨 / ⏭ 아직 안 해봄 / — 해당 없음")
r=line(r,"막혔을 때","무리해서 진행하지 말고, 그 화면을 캡처 → '문제 적기'에 한 줄 적기 → 결과는 ❌ 안 됨 으로 두고 다음으로 넘어가세요.")

r+=1
r=block(r,"④ 모르는 말 사전 (필요할 때만)")
gloss=[("회차","번역할 원문 한 편(에피소드 1화, 2화 같은 단위)."),
("설정집 / TB","인물 이름·용어·호칭 등을 정리한 엑셀 파일. AI가 이걸 보고 이름을 통일해서 번역해요."),
("프롬프트","AI에게 주는 '이렇게 번역해줘'라는 지시문. 1~5단계로 나뉘어 있어요."),
("모의(mock) 모드","진짜 AI를 부르지 않고 가짜 응답으로 흐름만 보는 연습 모드. 돈이 안 들어요. 내용은 자리표시자라 엉성한 게 정상."),
("Gemini","구글의 진짜 AI 모델. 여기서 실제 번역 품질을 봅니다."),
("평가표","번역이 잘됐는지 채점하려고 자동으로 만들어주는 엑셀 점수표. (한→영·한→일 작품에서만 생겨요.)"),
("STEP1~5","번역이 만들어지는 5개의 처리 단계. 화면에서 초록불로 하나씩 지나갑니다.")]
for k,v in gloss:
    r=line(r,k,v)

# ===== 2. 단계별 따라하기 =====
s=wb.create_sheet("단계별 따라하기"); s.sheet_view.showGridLines=False; s.freeze_panes="A2"
W={"A":8,"B":18,"C":48,"D":42,"E":8,"F":14,"G":26}
for c,w in W.items(): s.column_dimensions[c].width=w
heads=["순서","무엇을 점검","어떻게 해보나요","이렇게 나오면 정상","꼭 통과","결과","메모 / 문제"]
for i,h in enumerate(heads,1): hd(s.cell(1,i),h)
s.row_dimensions[1].height=26

rounds=[
("0","처음 준비","앱(MTPE.app) 더블클릭해서 화면 열기 → 번역 탭에서 모델 'Gemini' 고르고 키 붙여넣어 [저장] → 회차 원문 1~2개와 설정집(TB) 1개 올리기","화면이 뜨고, 키가 '설정됨 ✅', 올린 회차와 TB가 목록에 보임","⭐"),
("1","연습(모의)으로 전체 흐름","모델을 '🧪 키 없이 테스트(모의)'로 바꾸기 → 회차 1개 골라 번역 → 결과 [복사]/[다운로드] 눌러보기 → 프롬프트 관리 탭에서 [불러오기], [시험 실행]도 눌러보기","1~5단계가 초록불로 지나가고 결과가 뜸. 버튼이 다 눌림. (가짜 응답이라 내용이 엉성한 건 정상)","⭐"),
("2","진짜 번역(Gemini)","모델을 'Gemini'로 바꾸고 회차 1개를 실제로 번역","5단계가 진짜로 돌고, 최종 결과가 납품할 만큼 자연스러움(한→영이면 영어로, 중→한이면 한국어로). 설정집의 이름·용어가 쓰이고, 끝까지 잘림 없이 나옴","⭐"),
("2.5","평가표 만들기","(한→영·한→일 작품에서) 번역 끝난 뒤 [평가 시트 생성] 누르기","엑셀 평가표가 1부 생기고 원문·번역이 채워짐. 표 안의 계산(수식)이 안 깨짐. ※ 중→한 작품은 이 버튼이 없는 게 정상","⭐"),
("3","프롬프트 다듬기","프롬프트 관리 탭 → [시험 실행] → 마음에 안 드는 단계의 글을 고치고 [↻ 이 단계부터 다시] → 좋으면 [새 버전으로 저장]","고친 내용이 결과에 바로 반영되고, 저장한 새 버전이 번역 탭에서 최신으로 잡힘","⭐"),
("4","여러 회차 한꺼번에","회차 2~3개를 한 번에 골라 번역","하나씩 순서대로 처리되고, 회차마다 결과가 따로 나오고 각각 다운로드됨",""),
("5","일부러 이상하게 해보기","키를 틀리게/비우고 실행 · 회차 안 고르고 번역 눌러보기 · 아주 긴 원문 넣기 · 번역 도중 다른 탭 눌러보기","앱이 꺼지지 않고, 사람이 알아볼 안내 메시지가 뜸 (그냥 멈추거나 꺼지면 ❌)","⭐"),
("6","껐다 켜도 그대로인지","앱을 완전히 끄고 다시 켜기","고쳐둔 프롬프트·회차 목록·TB·번역 결과가 그대로 남아 있음","⭐"),
("7","윈도우에서도 (선택)","윈도우 PC가 있으면 MTPE.exe 실행해 순서 1~2만 가볍게 반복","맥에서와 똑같이 동작함",""),
]
core_src={}
for ri,(lbl,name,doo,ok,star) in enumerate(rounds,start=2):
    core=star=="⭐"; bg=CORE if core else None
    st(s.cell(ri,1),align="center",bold=core,bg=bg); s.cell(ri,1).value=lbl
    st(s.cell(ri,2),align="left",wrap=True,bg=bg); s.cell(ri,2).value=name
    st(s.cell(ri,3),align="left",wrap=True,bg=bg,top=True); s.cell(ri,3).value=doo
    st(s.cell(ri,4),align="left",wrap=True,bg=bg,top=True); s.cell(ri,4).value=ok
    st(s.cell(ri,5),align="center",bg=bg); s.cell(ri,5).value=star
    st(s.cell(ri,6),align="center",bold=True,bg=FILLIN)
    st(s.cell(ri,7),align="left",wrap=True,bg=FILLIN,top=True)
    s.row_dimensions[ri].height=72
    if core: core_src[lbl]=ri
dv=DataValidation(type="list",formula1='"✅ 잘 됨,❌ 안 됨,⏭ 아직 안 해봄,— 해당 없음"',allow_blank=True)
s.add_data_validation(dv); dv.add(f"F2:F{len(rounds)+1}")

# ===== 3. 세부 점검표 =====
ck=wb.create_sheet("세부 점검표"); ck.sheet_view.showGridLines=False; ck.freeze_panes="A3"
CW={"A":18,"B":5,"C":62,"D":13,"E":18,"F":28}
for c,w in CW.items(): ck.column_dimensions[c].width=w
ck.merge_cells("A1:F1")
st(ck["A1"],size=9,color="555555",align="left",bd=False)
ck["A1"]="자세히 확인하고 싶을 때 쓰는 보조 점검표(선택). '단계별 따라하기'만 해도 충분합니다."
for i,h in enumerate(["구분","#","확인할 내용","결과","증빙(캡처 파일명)","메모"],1): hd(ck.cell(2,i),h)
ck.row_dimensions[2].height=24

chk={
"실행 / 화면":[("앱(또는 서버)이 정상 실행되고 화면이 뜬다",),("위쪽 [번역]/[프롬프트 관리] 탭 전환이 된다",),("(윈도우) MTPE.exe 더블클릭하면 화면이 뜬다",)],
"모델 & 키":[("모델 목록에 Gemini/Claude/GPT/모의가 보인다",),("Gemini 고르고 키 저장하면 '설정됨 ✅'으로 바뀐다",),("키 없는 모델을 고르면 빨간 '미설정 ❌' 경고가 뜬다",),("🧪 모의 모델은 키 없이도 경고 없이 진행된다",),("(서버 공용키 모드) 키 입력칸이 숨고 '회사 공용 키' 안내가 뜬다",)],
"설정집(TB)":[("양식 템플릿 다운로드(.xlsx)가 받아지고 열린다",),("TB 파일 올리면 파일명이 초록색으로 표시된다",),("'AI에 보낼 시트 선택'이 그 파일의 진짜 시트만 보여준다",),("시트 체크가 저장된다(다시 들어가도 유지)",),("TB 삭제(✕)가 된다",)],
"회차 원문":[("폴더 경로 넣고 [적용]하면 그 폴더 회차가 목록에 뜬다",),("회차를 개별로 올리면 목록에 추가된다(txt/docx 등)",),("올린 회차는 ✕로 지워지고, 폴더 회차는 ✕가 없다(보호)",),("전체/해제로 회차를 한 번에 고를 수 있다",),("고른 회차 개수가 번역 버튼에 반영된다",)],
"번역 — 모의":[("회차 1개 번역하면 STEP1~5가 순서대로 초록불로 간다",),("각 단계에 '모의 응답' 결과가 뜬다",),("회차 여러 개를 고르면 순서대로 처리되고 결과가 따로 나온다",),("최종 결과 [복사]/[다운로드]가 동작한다",)],
"번역 — 진짜(Gemini) ⭐":[("Gemini로 회차 1개 번역하면 5단계 모두 진짜 결과가 나온다",),("STEP1 결과가 깔끔한 형태로 나온다",),("STEP3·STEP5 결과가 자연스러운 번역(목표 언어)으로 나온다",),("설정집의 이름·용어가 번역에 쓰인다",),("원문이 길어도 끝까지 잘림 없이 처리된다",),("결과 파일이 저장된다(works 폴더 안 final.txt)",),("(품질) 사람이 읽어 납품해도 될 수준인지 본다",)],
"프롬프트 관리":[("작품/언어/버전 [불러오기]하면 1~5단계 글이 채워진다",),("단계 글 수정 후 [현재 버전 덮어쓰기] 저장된다",),("[새 버전으로 저장]하면 버전이 늘고 번역 탭에서 최신으로 잡힌다",),("[이 버전 삭제]가 된다(확인창 후)",),("[+ 새 작품 만들기 (v1)]를 누르면 그 작품이 v1으로 즉시 생성되고 번역 탭 목록에도 바로 보인다",),("이미 있는 작품·언어를 다시 만들려 하면 경고가 뜬다(덮어쓰기 방지)",)],
"프롬프트 다듬기 ⭐":[("샘플 회차+모델 고르고 [시험 실행]하면 단계별 결과가 뜬다",),("Gemini로 시험 실행하면 진짜 결과가 단계별로 나온다",),("글 고치고 [이 단계부터 다시] 누르면 그 단계만 다시 돌고 반영된다",),("[실제 보낸 프롬프트 보기]로 합쳐진 지시문이 보인다",),("다듬은 뒤 저장하면 번역 탭에 반영된다",)],
"이상 상황 / 안 죽는지":[("틀린/만료된 키로 하면 알아볼 에러 메시지가 뜬다(안 꺼짐)",),("회차 안 고르고 번역하면 안내 메시지가 뜬다",),("아주 긴 원문/빈 파일을 넣어도 어떻게든 처리/안내된다",),("번역 도중 다른 탭으로 옮기거나 다시 해도 안 깨진다",)],
"껐다 켜도 유지":[("앱을 껐다 켜도 고쳐둔 프롬프트가 남아 있다",),("회차 목록·TB·시트 선택이 다시 켜도 유지된다",),("번역 결과(final.txt)가 보관된다",)],
"평가표 ⭐ (한→영·한→일)":[("번역 후 [평가 시트 생성] 버튼이 보인다",),("(중→한 작품) 평가표 버튼이 안 보인다(정상)",),("[평가 시트 생성]하면 엑셀이 생기고 다운로드/폴더 열기가 된다",),("평가표에 원문·번역이 채워지고 표 안 계산이 안 깨진다",),("장면 설명, 버전·날짜 등이 표시된다",),("샘플 문장이 채워지고 채점칸은 비어 있다(직접 채점용)",),("[평가 시트 생성]하면 평가표가 1부만 생성된다",)],
}
dv2=DataValidation(type="list",formula1='"✅ 잘 됨,❌ 안 됨,⏭ 아직 안 해봄,— 해당 없음"',allow_blank=True)
ck.add_data_validation(dv2)
row=3
for sec,items in chk.items():
    core="⭐" in sec
    for j,(it,) in enumerate(items,1):
        bg=CORE if core else None
        if j==1:
            st(ck.cell(row,1),bold=True,align="left",wrap=True,bg=SECT); ck.cell(row,1).value=sec
        else:
            st(ck.cell(row,1),bg=SECT)
        st(ck.cell(row,2),align="center",bg=bg); ck.cell(row,2).value=j
        st(ck.cell(row,3),align="left",wrap=True,bg=bg); ck.cell(row,3).value=it
        st(ck.cell(row,4),align="center",bold=True,bg=FILLIN)
        st(ck.cell(row,5),align="left",bg=FILLIN)
        st(ck.cell(row,6),align="left",wrap=True,bg=FILLIN)
        ck.row_dimensions[row].height=26
        row+=1
dv2.add(f"D3:D{row-1}")
row=3
for sec,items in chk.items():
    n=len(items)
    if n>1:
        ck.merge_cells(start_row=row,start_column=1,end_row=row+n-1,end_column=1)
        ck.cell(row,1).alignment=Alignment(horizontal="left",vertical="center",wrap_text=True)
    row+=n

# ===== 4. 문제 적기 =====
ig=wb.create_sheet("문제 적기"); ig.sheet_view.showGridLines=False; ig.freeze_panes="A3"
IW={"A":5,"B":12,"C":24,"D":40,"E":34,"F":12,"G":14,"H":22}
for c,w in IW.items(): ig.column_dimensions[c].width=w
ig.merge_cells("A1:H1")
st(ig["A1"],size=9,color="555555",align="left",bd=False)
ig["A1"]="하다가 이상하거나 막히면 여기 적으세요. 화면을 캡처해 두고 파일명을 메모에 함께 적으면 전달하기 좋아요."
for i,h in enumerate(["#","발견일","어디서 (화면/순서)","무슨 문제인가요","어떻게 하면 다시 나오나요","급한 정도","처리 상태","메모"],1):
    hd(ig.cell(2,i),h)
ig.row_dimensions[2].height=24
for ri in range(3,28):
    ig.cell(ri,1).value=ri-2
    for ci in range(1,9):
        st(ig.cell(ri,ci),align="center" if ci in(1,2,6,7) else "left",wrap=ci in(4,5,8),bg=FILLIN if ci!=1 else LGRAY)
    ig.row_dimensions[ri].height=30
dvs=DataValidation(type="list",formula1='"🔴 막힘(출시 불가),🟡 불편(고쳐야 함),🟢 사소(개선)"',allow_blank=True)
dvt=DataValidation(type="list",formula1='"새로 발견,고치는 중,고침 완료,다시 해보니 괜찮음,나중에"',allow_blank=True)
ig.add_data_validation(dvs); ig.add_data_validation(dvt)
dvs.add("F3:F27"); dvt.add("G3:G27")
ig.merge_cells("A29:H29")
st(ig["A29"],size=9,color="555555",align="left",bd=False)
ig["A29"]="급한 정도 — 🔴 막힘: 핵심 기능을 못 씀(출시 불가) · 🟡 불편: 쓸 수는 있지만 고쳐야 함 · 🟢 사소: 있으면 좋은 개선. '🔴 막힘'이 안 고쳐진 채 남으면 종합 결과가 자동으로 '아직 출시 안 됨'으로 표시됩니다."

# ===== 5. 종합 결과 =====
v=wb.create_sheet("종합 결과"); v.sheet_view.showGridLines=False
v.column_dimensions["A"].width=3
v.column_dimensions["B"].width=24
v.column_dimensions["C"].width=30
v.column_dimensions["D"].width=18
v.column_dimensions["E"].width=16
v.merge_cells("B2:E2"); st(v["B2"],bold=True,size=16,color="FFFFFF",bg=NAVY)
v["B2"]="  종합 결과 (자동으로 계산됩니다)"
v.row_dimensions[2].height=30
v.merge_cells("B3:E3"); st(v["B3"],size=9,color="FFFFFF",bg=TEAL,align="left")
v["B3"]="  '단계별 따라하기'와 '문제 적기'를 채우면 아래가 저절로 바뀝니다. 직접 계산할 필요 없어요."

r=5
v.merge_cells(f"B{r}:E{r}"); st(v[f"B{r}"],bold=True,size=11,color="FFFFFF",bg=TEAL,align="left")
v[f"B{r}"]="  ① 꼭 통과해야 하는 항목 (이게 다 '잘 됨'이어야 출시)"; r+=1
hd(v.cell(r,2),"순서"); v.merge_cells(f"C{r}:D{r}"); hd(v.cell(r,3),"무엇"); hd(v.cell(r,5),"결과"); r+=1
gate=[("0","처음 준비"),("1","연습(모의) 흐름"),("2","진짜 번역"),("2.5","평가표 만들기"),
      ("3","프롬프트 다듬기"),("5","이상 상황"),("6","껐다 켜도 유지")]
judge_cells=[]
for lbl,name in gate:
    src=core_src[lbl]
    st(v.cell(r,2),align="center"); v.cell(r,2).value=lbl
    v.merge_cells(f"C{r}:D{r}"); st(v.cell(r,3),align="left"); v.cell(r,3).value=name
    cell=f"'단계별 따라하기'!F{src}"
    st(v.cell(r,5),align="center",bold=True)
    v.cell(r,5).value=f'=IF({cell}="","– 아직",{cell})'
    judge_cells.append(cell)
    r+=1

r+=1
v.merge_cells(f"B{r}:E{r}"); st(v[f"B{r}"],bold=True,size=11,color="FFFFFF",bg=TEAL,align="left")
v[f"B{r}"]="  ② 최종 판단"; r+=1
st(v.cell(r,2),bold=True,bg=LGRAY,align="left"); v.cell(r,2).value="안 고쳐진 '🔴 막힘' 문제 수"
v.merge_cells(f"C{r}:D{r}")
miss=('=COUNTIFS(\'문제 적기\'!F:F,"🔴 막힘(출시 불가)",\'문제 적기\'!G:G,"새로 발견")'
      '+COUNTIFS(\'문제 적기\'!F:F,"🔴 막힘(출시 불가)",\'문제 적기\'!G:G,"고치는 중")'
      '+COUNTIFS(\'문제 적기\'!F:F,"🔴 막힘(출시 불가)",\'문제 적기\'!G:G,"나중에")')
st(v.cell(r,3),align="center",bold=True); v.cell(r,3).value=miss
miss_cell=f"C{r}"; r+=1
st(v.cell(r,2),bold=True,bg=LGRAY,align="left"); v.cell(r,2).value="자동 판정"
v.merge_cells(f"C{r}:E{r}")
allp="AND("+",".join(f'{c}="✅ 잘 됨"' for c in judge_cells)+")"
st(v.cell(r,3),bold=True,size=13,align="center",bg=FILLIN)
v.cell(r,3).value=f'=IF(AND({allp},{miss_cell}=0),"🟢 출시해도 됩니다","🔴 아직 출시 안 됨 (위 항목/막힘 문제 확인)")'
v.row_dimensions[r].height=24; r+=1
st(v.cell(r,2),bold=True,bg=LGRAY,align="left"); v.cell(r,2).value="내 최종 결정"
v.merge_cells(f"C{r}:E{r}"); st(v.cell(r,3),bg=FILLIN,align="left")
v.cell(r,3).value="☐ 출시   ☐ 문제 적고 조건부   ☐ 보류"
r+=2
v.merge_cells(f"B{r}:E{r}"); st(v[f"B{r}"],size=9,color="555555",align="left",bd=False,wrap=True)
v[f"B{r}"]="자동 판정 규칙: 위 ①의 7개 항목이 모두 '✅ 잘 됨'이고, '문제 적기'에 안 고쳐진 🔴 막힘 문제가 0개일 때만 '출시해도 됩니다'로 바뀝니다. (순서 4·7은 선택이라 판정에 안 들어갑니다.)"
v.row_dimensions[r].height=40

import os
_out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MTPE_QC_기록부.xlsx")
wb.save(_out)
print("ok:", _out)
