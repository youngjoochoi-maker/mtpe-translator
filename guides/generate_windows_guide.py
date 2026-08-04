"""
MTPE 윈도우 — 다운로드·설치·운영 가이드(엑셀) 생성기

사용법:
    pip install openpyxl
    python generate_windows_guide.py
  → 같은 폴더에 MTPE_윈도우_설치운영_가이드.xlsx 생성

개발 경험이 없는 분도 따라 할 수 있게, '받기 → 실행 → 운영' 순서로 정리한 안내서입니다.
프로덕트가 바뀌면(다운로드 위치·기능 등) 아래 sheet 내용만 고쳐 다시 실행하면 됩니다.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# 저장소에 다른 프로그램(NovelSplitter·api) 릴리스도 있어 /releases/latest 는 MTPE 가 아닐 수 있음.
# → 전체 릴리스 목록에서 vX.Y.Z (MTPE.exe 첨부) 중 최신을 받는다.
RELEASE_URL = "https://github.com/youngjoochoi-maker/mtpe-translator/releases"

FONT = "맑은 고딕"
NAVY = "1F4E5F"; TEAL = "2E7D8A"; LGRAY = "EEF3F5"; SECT = "D7E6EA"
CORE = "FFF3D6"; FILLIN = "FFFBEA"; WARN = "FBE3E0"
thin = Side(style="thin", color="BFCBD0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def st(c, *, bold=False, size=10, color="000000", bg=None, align="left", wrap=False, bd=True, top=False):
    c.font = Font(name=FONT, bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="top" if top else "center", wrap_text=wrap)
    if bg: c.fill = PatternFill("solid", fgColor=bg)
    if bd: c.border = border
    return c


def hd(c, t, bg=NAVY):
    c.value = t; st(c, bold=True, color="FFFFFF", bg=bg, align="center", wrap=True)


wb = Workbook()

# ============================================================
# 1. 시작하기
# ============================================================
g = wb.active; g.title = "시작하기"; g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 3
g.column_dimensions["B"].width = 22
g.column_dimensions["C"].width = 88
g.merge_cells("B2:C2"); st(g["B2"], bold=True, size=18, color="FFFFFF", bg=NAVY)
g["B2"] = "  MTPE 윈도우 — 다운로드 · 설치 · 운영 가이드 (처음 하는 분용)"
g.row_dimensions[2].height = 34


def block(row, title):
    g.merge_cells(f"B{row}:C{row}")
    st(g[f"B{row}"], bold=True, size=11, color="FFFFFF", bg=TEAL); g[f"B{row}"] = "  " + title
    return row + 1


def line(row, k, v, kbg=LGRAY):
    st(g[f"B{row}"], bold=True, bg=kbg, align="left", wrap=True); g[f"B{row}"] = k
    st(g[f"C{row}"], align="left", wrap=True); g[f"C{row}"] = v
    g.row_dimensions[row].height = max(16, 18 * (len(v) // 54 + 1))
    return row + 1


r = 4
r = block(r, "① 이 프로그램은 무엇인가요?")
r = line(r, "한 줄 설명", "MTPE는 웹소설 원문을 5단계로 처리해 기계번역(MT)을 뽑아내는 프로그램입니다. 윈도우에서는 설치 없이 파일 하나(MTPE.exe)만 더블클릭하면 바로 켜집니다.")
r = line(r, "무엇을 하나요", "작품·언어를 고르고, 설정집(인물·용어 엑셀)과 회차 원문을 넣으면 → 번역 결과를 만들어 줍니다. (한→영·한→일 작품은 채점용 '평가표' 엑셀도 자동 생성)")
r = line(r, "개발 지식 필요?", "전혀 필요 없습니다. 화면을 눌러서 쓰는 일반 프로그램이에요. 이 가이드의 시트를 위에서 아래로 따라 하면 됩니다.")

r += 1
r = block(r, "② 딱 3단계예요")
r = line(r, "1단계 — 받기", "MTPE.exe 파일을 받습니다. (가장 쉬운 건 운영자에게서 파일을 받는 것 → '1. 다운로드 받기' 시트)")
r = line(r, "2단계 — 실행", "받은 MTPE.exe를 더블클릭. 설치 과정 없이 바로 켜집니다. (→ '2. 설치·첫 실행' 시트)")
r = line(r, "3단계 — 운영", "모델·키를 한 번 넣고, 작품·회차를 골라 번역. (→ '3. 운영(쓰기)' 시트)")

r += 1
r = block(r, "③ 준비물")
r = line(r, "PC", "윈도우 10 또는 11 (일반 사무용 PC면 충분)")
r = line(r, "인터넷", "필요 (AI 모델을 인터넷으로 부릅니다)")
r = line(r, "API 키", "실제 번역에는 AI 키가 필요합니다. 본인 키를 넣거나, 회사 공용 방식을 쓰면 됩니다. 키 없이 '연습(모의) 모드'로 화면만 먼저 볼 수도 있어요.")

r += 1
r = block(r, "④ 어느 시트부터 보나요?")
r = line(r, "순서대로", "'1. 다운로드 받기' → '2. 설치·첫 실행' → '3. 운영(쓰기)'. 막히면 '4. 문제 해결'을 보세요. 끝나면 '5. 설치 확인'에 체크.")

# ============================================================
# 2. 다운로드 받기
# ============================================================
d = wb.create_sheet("1. 다운로드 받기"); d.sheet_view.showGridLines = False; d.freeze_panes = "A2"
DW = {"A": 8, "B": 26, "C": 78, "D": 30}
for c, w in DW.items(): d.column_dimensions[c].width = w
for i, h in enumerate(["순서", "할 일", "자세히", "이렇게 되면 정상"], 1): hd(d.cell(1, i), h)
d.row_dimensions[1].height = 26


def drow(row, n, todo, detail, ok, bg=None):
    st(d.cell(row, 1), align="center", bold=True, bg=bg); d.cell(row, 1).value = n
    st(d.cell(row, 2), align="left", wrap=True, bg=bg, top=True); d.cell(row, 2).value = todo
    st(d.cell(row, 3), align="left", wrap=True, bg=bg, top=True); d.cell(row, 3).value = detail
    st(d.cell(row, 4), align="left", wrap=True, bg=bg, top=True); d.cell(row, 4).value = ok
    d.row_dimensions[row].height = max(40, 16 * (len(detail) // 50 + 1))
    return row + 1


rr = 2
d.merge_cells(f"A{rr}:D{rr}"); st(d.cell(rr, 1), bold=True, color="FFFFFF", bg=TEAL, align="left")
d.cell(rr, 1).value = "  방법 A — 운영자에게 파일 받기 (가장 쉬움 · 개발지식·GitHub 계정 불필요) ⭐추천"; rr += 1
rr = drow(rr, "A-1", "운영자에게 MTPE.exe 요청", "회사에서 MTPE를 관리하는 담당자(운영자)에게 'MTPE.exe 파일 주세요'라고 요청합니다. 사내 공유 드라이브·메신저·USB 등 어떤 방법으로든 파일만 받으면 됩니다.", "MTPE.exe 파일을 받음 (약 58MB)", bg=CORE)
rr = drow(rr, "A-2", "내 PC에 저장", "받은 MTPE.exe를 바탕화면이나 '문서' 폴더 등 찾기 쉬운 곳에 둡니다. 압축(.zip)으로 받았으면 마우스 오른쪽 → '압축 풀기' 후 MTPE.exe를 꺼냅니다.", "바탕화면 등에 MTPE.exe가 보임", bg=CORE)

rr += 1
d.merge_cells(f"A{rr}:D{rr}"); st(d.cell(rr, 1), bold=True, color="FFFFFF", bg=TEAL, align="left")
d.cell(rr, 1).value = "  방법 B — GitHub에서 직접 받기 (GitHub 계정 + 저장소 접근 권한 필요)"; rr += 1
rr = drow(rr, "B-1", "GitHub 로그인", "인터넷 브라우저에서 github.com 에 로그인합니다. (계정이 없으면 방법 A로 받으세요. 이 저장소는 비공개라 권한이 있어야 보입니다.)", "GitHub에 로그인된 상태")
rr = drow(rr, "B-2", "릴리스 목록 열기", f"주소창에 아래를 입력해 엽니다:\n{RELEASE_URL}\n(권한이 없으면 'Not Found'가 떠요 → 운영자에게 '저장소에 초대(collaborator)해 주세요' 요청)", "여러 버전(Releases) 목록이 보임")
rr = drow(rr, "B-3", "MTPE 최신 버전 고르기", "목록에서 이름이 v1.1.1 처럼 v로 시작하는 것 중 가장 위(최신)를 클릭. ※ splitter-… / api-… 로 시작하는 건 다른 프로그램이니 무시하세요.", "MTPE.exe 가 첨부된 버전 화면이 열림")
rr = drow(rr, "B-4", "MTPE.exe 내려받기", "그 화면 아래 'Assets' 목록에서 MTPE.exe 를 클릭하면 다운로드됩니다.", "다운로드 폴더에 MTPE.exe (약 58MB)")

rr += 1
d.merge_cells(f"A{rr}:D{rr}")
st(d.cell(rr, 1), size=9, color="555555", align="left", wrap=True, bd=False)
d.cell(rr, 1).value = ("※ 비개발자라면 방법 A(운영자에게 파일 받기)가 제일 쉽습니다. 운영자 한 명이 방법 B로 한 번 받아두고, "
                       "그 MTPE.exe를 사내에 공유하면 나머지 분들은 받기만 하면 됩니다. ※ 프로그램이 업데이트되면 새 MTPE.exe로 같은 방식으로 교체하면 됩니다.")
d.row_dimensions[rr].height = 44

# ============================================================
# 3. 설치 · 첫 실행
# ============================================================
s = wb.create_sheet("2. 설치·첫 실행"); s.sheet_view.showGridLines = False; s.freeze_panes = "A2"
SW = {"A": 8, "B": 26, "C": 78, "D": 30}
for c, w in SW.items(): s.column_dimensions[c].width = w
for i, h in enumerate(["순서", "할 일", "자세히", "이렇게 되면 정상"], 1): hd(s.cell(1, i), h)
s.row_dimensions[1].height = 26


def srow(row, n, todo, detail, ok, bg=None):
    st(s.cell(row, 1), align="center", bold=True, bg=bg); s.cell(row, 1).value = n
    st(s.cell(row, 2), align="left", wrap=True, bg=bg, top=True); s.cell(row, 2).value = todo
    st(s.cell(row, 3), align="left", wrap=True, bg=bg, top=True); s.cell(row, 3).value = detail
    st(s.cell(row, 4), align="left", wrap=True, bg=bg, top=True); s.cell(row, 4).value = ok
    s.row_dimensions[row].height = max(44, 16 * (len(detail) // 50 + 1))
    return row + 1


rr = 2
rr = srow(rr, "1", "설치 필요 없음 — 더블클릭", "MTPE.exe를 더블클릭하면 그게 곧 실행입니다. 따로 설치(Install) 과정이 없습니다.", "잠시 후 프로그램 창이 뜸")
rr = srow(rr, "2", "파란 경고가 뜨면 (중요)", "'Windows의 PC를 보호했습니다'(SmartScreen) 파란 창이 뜰 수 있어요. → '추가 정보' 클릭 → '실행' 버튼 클릭. (사내용이라 서명이 없어서 뜨는 정상 경고입니다.)", "'추가 정보 → 실행'을 누르면 프로그램이 켜짐", bg=WARN)
rr = srow(rr, "3", "첫 실행은 조금 느림", "처음 켤 때 파일을 푸느라 몇 초~십여 초 걸립니다. 검은 창이 잠깐 떴다 사라질 수 있어요(정상).", "잠시 기다리면 프로그램 화면이 뜸")
rr = srow(rr, "4", "화면 확인", "프로그램이 자체 창으로 열립니다. 만약 창 대신 인터넷 브라우저로 열려도 정상이에요(같은 화면).", "위쪽에 [번역] / [프롬프트 관리] 탭이 보임")
rr = srow(rr, "5", "모델 + 키 넣기 (첫 1회)", "[번역] 탭의 '1. 모델 & API 키'에서 모델을 고르고 API 키를 붙여넣어 [저장]. → '설정됨 ✅' 확인. (회사 공용 방식이면 운영자 안내를 따르세요.)", "키가 '설정됨 ✅'으로 바뀜", bg=CORE)
rr = srow(rr, "6", "키 없이 먼저 보려면", "모델에서 '🧪 키 없이 테스트(모의)'를 고르면 키 없이 화면 흐름만 연습할 수 있어요(결과는 가짜).", "모의 모델로 번역을 눌러도 경고 없이 진행됨")
rr = srow(rr, "7", "끄기", "프로그램 창을 닫으면 종료됩니다. (브라우저로 열렸으면 같이 뜬 검은 창도 닫기)", "창을 닫으면 프로그램이 꺼짐")

rr += 1
s.merge_cells(f"A{rr}:D{rr}")
st(s.cell(rr, 1), size=9, color="555555", align="left", wrap=True, bd=False)
s.cell(rr, 1).value = ("※ '바로 꺼져요/안 떠요' → '4. 문제 해결' 시트를 보세요. ※ 키는 이 PC 안에만 저장됩니다(외부 전송 없음). "
                       "저장 위치: 윈도우 탐색기 주소창에 %APPDATA%\\MTPE 입력.")
s.row_dimensions[rr].height = 32

# ============================================================
# 4. 운영(쓰기)
# ============================================================
o = wb.create_sheet("3. 운영(쓰기)"); o.sheet_view.showGridLines = False; o.freeze_panes = "A2"
OW = {"A": 8, "B": 26, "C": 86, "D": 30}
for c, w in OW.items(): o.column_dimensions[c].width = w
for i, h in enumerate(["순서", "무엇", "어떻게", "결과"], 1): hd(o.cell(1, i), h)
o.row_dimensions[1].height = 26


def orow(row, n, todo, detail, ok, bg=None):
    st(o.cell(row, 1), align="center", bold=True, bg=bg); o.cell(row, 1).value = n
    st(o.cell(row, 2), align="left", wrap=True, bg=bg, top=True); o.cell(row, 2).value = todo
    st(o.cell(row, 3), align="left", wrap=True, bg=bg, top=True); o.cell(row, 3).value = detail
    st(o.cell(row, 4), align="left", wrap=True, bg=bg, top=True); o.cell(row, 4).value = ok
    o.row_dimensions[row].height = max(46, 16 * (len(detail) // 56 + 1))
    return row + 1


rr = 2
rr = orow(rr, "1", "작품·언어 고르기", "[번역] 탭에서 모델, 그리고 번역할 작품 / 언어 / 버전(기본 최신)을 선택합니다.", "선택한 작품·언어가 표시됨")
rr = orow(rr, "2", "설정집(TB) 올리기", "인물·용어 엑셀(.xlsx)을 [업로드/교체]로 올리고, 'AI에 보낼 시트 선택'에서 쓸 시트를 체크. 양식이 필요하면 '양식 템플릿 내려받기'로 받아 작성.", "파일명이 초록색으로 뜨고 시트 체크가 보임")
rr = orow(rr, "3", "회차 원문 넣기", "폴더 경로를 넣고 [적용]하거나, 파일을 [회차 추가]로 올립니다(txt/docx 등). 번역할 회차를 체크.", "회차 목록에 뜨고 선택 개수가 버튼에 반영")
rr = orow(rr, "4", "번역 실행", "[▶ 선택 회차 번역] 클릭 → 오른쪽에 STEP1~5가 순서대로 진행되고 회차별 결과가 나옴.", "5단계 초록불 → 최종 결과 표시", bg=CORE)
rr = orow(rr, "5", "결과 가져가기", "최종 결과 [복사] 또는 [다운로드]. 자동으로도 저장됩니다(아래 '파일 위치' 참고).", "결과를 복사/다운로드함")
rr = orow(rr, "6", "평가표 만들기 (한→영·한→일)", "번역 후 최종 결과에서 [평가 시트 생성] → 채점용 엑셀 평가표가 1부 생성. (중→한 작품은 이 버튼이 없는 게 정상)", "평가표 .xlsx 1개 생성·다운로드", bg=CORE)
rr = orow(rr, "7", "프롬프트 고치기 (담당자)", "[프롬프트 관리] 탭에서 단계별 지시문을 불러와 수정·저장하거나 새 버전으로 저장. [시험 실행]으로 미리 테스트.", "수정·새 버전이 번역 탭에 반영됨")

rr += 1
o.merge_cells(f"A{rr}:D{rr}"); st(o.cell(rr, 1), bold=True, color="FFFFFF", bg=TEAL, align="left")
o.cell(rr, 1).value = "  파일은 어디에 저장되나요?"; rr += 1


def opath(row, k, v):
    st(o.cell(row, 2), bold=True, bg=LGRAY, align="left", wrap=True); o.cell(row, 2).value = k
    o.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    st(o.cell(row, 3), align="left", wrap=True); o.cell(row, 3).value = v
    o.row_dimensions[row].height = 26
    return row + 1


rr = opath(rr, "데이터 폴더", "윈도우 탐색기 주소창에  %APPDATA%\\MTPE  입력")
rr = opath(rr, "번역 결과", "works\\<작품>\\output\\<회차>\\final.txt")
rr = opath(rr, "평가표", "works\\<작품>\\eval\\<회차>\\<작품>_<버전>_평가표.xlsx")
rr = opath(rr, "API 키", ".env  (이 PC에만 저장)")

rr += 1
o.merge_cells(f"A{rr}:D{rr}")
st(o.cell(rr, 1), size=9, color="555555", align="left", wrap=True, bd=False)
o.cell(rr, 1).value = ("※ 여러 명이 '회사 공용 키' 하나로 쓰려면, 각자 PC에 .exe를 까는 것보다 '서버 배포'(브라우저로 접속)가 더 적합합니다. "
                       "운영자에게 문의하세요. ※ 평가표는 1부만 나옵니다 — 2명 이상 평가는 운영측이 파일을 복제·이름변경해 나눠 줍니다.")
o.row_dimensions[rr].height = 36

# ============================================================
# 5. 문제 해결
# ============================================================
f = wb.create_sheet("4. 문제 해결"); f.sheet_view.showGridLines = False; f.freeze_panes = "A2"
FW = {"A": 30, "B": 64, "C": 44}
for c, w in FW.items(): f.column_dimensions[c].width = w
for i, h in enumerate(["이런 일이 생기면", "왜 그런가요", "이렇게 하세요"], 1): hd(f.cell(1, i), h)
f.row_dimensions[1].height = 26

faqs = [
    ("파란 경고창이 떠서 실행이 안 돼요", "서명 안 된 사내 프로그램이라 윈도우(SmartScreen)가 한 번 막는 정상 경고예요.", "'추가 정보' → '실행' 클릭. (악성 아님)"),
    ("더블클릭해도 창이 안 떠요 / 바로 꺼져요", "WebView2(엣지 구성요소)가 없거나, 백신이 잠시 막았을 수 있어요.", "잠깐 기다려 보세요(첫 실행 느림). 그래도 안 되면 운영자에게 문의 — 콘솔 메시지를 보는 디버그 빌드로 원인을 잡을 수 있어요."),
    ("창 대신 인터넷 브라우저로 열려요", "전용 창(WebView2)이 없으면 자동으로 브라우저로 여는 정상 동작이에요.", "그대로 쓰면 됩니다. 같은 화면이에요."),
    ("처음 켤 때 너무 느려요", "단일 파일이라 처음에 압축을 임시폴더에 푸는 시간이 필요해요.", "몇 초~십여 초 기다리면 떠요(정상). 두 번째부터는 빨라요."),
    ("키가 빨간 '미설정 ❌'으로 떠요", "그 모델에 맞는 API 키가 저장 안 된 상태예요.", "키를 저장하거나, 모델을 '🧪 모의'로 바꾸세요."),
    ("회차 목록에 엉뚱한 파일이 떠요", "원문 폴더에 원문이 아닌 파일이 섞여 있어요.", "회차 원문만 따로 폴더에 모아 경로를 지정하세요."),
    ("평가표 [평가 시트 생성] 버튼이 없어요", "평가표는 한→영·한→일 작품에만 있어요(중→한은 템플릿 없음).", "한→영·한→일 작품에서 번역 후 확인하세요. (정상)"),
    ("내가 만든 결과·설정이 사라졌어요", "데이터는 %APPDATA%\\MTPE 에 남아 있어요.", "탐색기 주소창에 %APPDATA%\\MTPE 입력해 prompts·works 확인."),
    ("새 버전으로 바꾸고 싶어요", "프로그램이 업데이트되면 새 MTPE.exe가 나와요.", "운영자에게 새 MTPE.exe를 받아 기존 파일과 교체(데이터는 그대로 유지)."),
]
rr = 2
for q, why, how in faqs:
    st(f.cell(rr, 1), bold=True, align="left", wrap=True, bg=SECT, top=True); f.cell(rr, 1).value = q
    st(f.cell(rr, 2), align="left", wrap=True, top=True); f.cell(rr, 2).value = why
    st(f.cell(rr, 3), align="left", wrap=True, bg=FILLIN, top=True); f.cell(rr, 3).value = how
    f.row_dimensions[rr].height = max(34, 16 * (len(why) // 40 + 1))
    rr += 1

# ============================================================
# 6. 설치 확인 (체크)
# ============================================================
ck = wb.create_sheet("5. 설치 확인"); ck.sheet_view.showGridLines = False; ck.freeze_panes = "A3"
CW = {"A": 6, "B": 64, "C": 16, "D": 40}
for c, w in CW.items(): ck.column_dimensions[c].width = w
ck.merge_cells("A1:D1")
st(ck["A1"], size=9, color="555555", align="left", bd=False)
ck["A1"] = "설치가 잘 됐는지 확인하는 짧은 체크표. 노란 칸을 누르면 ▼로 골라 넣을 수 있어요."
for i, h in enumerate(["#", "확인할 내용", "결과", "메모"], 1): hd(ck.cell(2, i), h)
ck.row_dimensions[2].height = 24

checks = [
    "MTPE.exe 파일을 받았다 (약 58MB)",
    "더블클릭하니 (경고 시 '추가 정보 → 실행' 후) 화면이 떴다",
    "위쪽 [번역] / [프롬프트 관리] 탭이 보인다",
    "모델 + API 키를 저장해 '설정됨 ✅'이 됐다 (또는 🧪 모의로 진행)",
    "🧪 모의 모드로 회차 1개 번역이 5단계로 돌았다",
    "(실제 키) Gemini 등으로 실제 번역 1회가 됐다",
    "(한→영·한→일) [평가 시트 생성]으로 평가표 1부가 만들어졌다",
    "껐다 켜도 내 설정·결과가 남아 있다",
]
dv = DataValidation(type="list", formula1='"✅ 됨,❌ 안 됨,⏭ 아직 안 함,— 해당 없음"', allow_blank=True)
ck.add_data_validation(dv)
rr = 3
for i, t in enumerate(checks, 1):
    st(ck.cell(rr, 1), align="center"); ck.cell(rr, 1).value = i
    st(ck.cell(rr, 2), align="left", wrap=True); ck.cell(rr, 2).value = t
    st(ck.cell(rr, 3), align="center", bold=True, bg=FILLIN)
    st(ck.cell(rr, 4), align="left", wrap=True, bg=FILLIN)
    ck.row_dimensions[rr].height = 28
    rr += 1
dv.add(f"C3:C{rr-1}")
ck.merge_cells(f"A{rr+1}:D{rr+1}")
st(ck.cell(rr + 1, 1), size=9, color="555555", align="left", wrap=True, bd=False)
ck.cell(rr + 1, 1).value = "위 1~5번(또는 6번)이 '✅ 됨'이면 설치·기본 동작 정상입니다. 막히면 '4. 문제 해결' 시트를 보세요."
ck.row_dimensions[rr + 1].height = 28

# ============================================================
_out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MTPE_윈도우_설치운영_가이드.xlsx")
wb.save(_out)
print("ok:", _out)
